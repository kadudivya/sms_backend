"""Read .xlsx, .xls and .csv uploads into rows (bulk uploads and data scripts)."""
import csv
import io
import os

from app.core.errors import ValidationError


def cell(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    v = str(v).strip()
    return v or None


def sheets(data, filename):
    """Yield each sheet as a list of row lists (raw values)."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        yield os.path.basename(filename), list(csv.reader(io.StringIO(text)))
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            yield ws.title, [list(r) for r in ws.iter_rows(values_only=True)]
    elif name.endswith(".xls"):
        try:
            import xlrd
        except ImportError as exc:  # pragma: no cover
            raise ValidationError("Install xlrd to read .xls files, or save the file as .xlsx") from exc
        book = xlrd.open_workbook(file_contents=data)
        for sh in book.sheets():
            yield sh.name, [sh.row_values(i) for i in range(sh.nrows)]
    else:
        raise ValidationError("Upload an .xlsx, .xls or .csv file", error_code="UNSUPPORTED_FILE_TYPE")


def read_rows(file_storage):
    """Plain table reader (first row = headers) for other bulk uploads."""
    data = file_storage.read()
    out = []
    for _, raw in sheets(data, file_storage.filename):
        if not raw:
            continue
        headers = [str(h).strip() if h is not None else "" for h in raw[0]]
        out += [{headers[i]: v for i, v in enumerate(r) if i < len(headers)} for r in raw[1:]
                if any(v not in (None, "") for v in r)]
        break
    return out
